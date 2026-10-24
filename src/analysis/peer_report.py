```python
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import os

from src.database.database import Database


class PeerReport:

    def __init__(self):

        self.db = Database()
        self.db.connect()

        self.output_folder = "output"

        os.makedirs(
            self.output_folder,
            exist_ok=True
        )

        self.workbook = Workbook()

        self.workbook.remove(
            self.workbook.active
        )

        self.green = PatternFill(
            fill_type="solid",
            start_color="90EE90"
        )

        self.yellow = PatternFill(
            fill_type="solid",
            start_color="FFF59D"
        )

        self.red = PatternFill(
            fill_type="solid",
            start_color="FFC7CE"
        )

        self.gold = PatternFill(
            fill_type="solid",
            start_color="FFD966"
        )

    def load_peer_groups(self):

        peer_groups = self.db.get_peer_groups()

        grouped = {}

        for group_name, company_id, benchmark in peer_groups:

            if group_name not in grouped:
                grouped[group_name] = []

            grouped[group_name].append(
                {
                    "company_id": company_id,
                    "benchmark": benchmark
                }
            )

        return grouped

    def create_sheets(self):

        groups = self.load_peer_groups()

        headers = [

            "Company ID",
            "Company Name",

            "Revenue Growth",
            "Profit Growth",

            "Net Profit Margin",
            "Operating Margin",

            "ROE",
            "ROCE",

            "Debt To Equity",
            "Interest Coverage",

            "Asset Turnover",

            "Free Cash Flow",

            "CFO Quality",

            "Capex Intensity",

            "Composite Score",

            "Revenue %",
            "Profit %",
            "NPM %",
            "ROE %",
            "ROCE %",
            "Composite %"

        ]

        for group_name in groups:

            sheet = self.workbook.create_sheet(
                title=group_name[:31]
            )

            for col, header in enumerate(headers, start=1):

                sheet.cell(
                    row=1,
                    column=col
                ).value = header

        print(f"{len(groups)} Sheets Created")

    def save(self):

        filename = os.path.join(
            self.output_folder,
            "peer_comparison.xlsx"
        )

        self.workbook.save(filename)

        print("=" * 60)
        print("Excel Report Saved")
        print(filename)
        print("=" * 60)

    def close(self):

        self.db.close()