import math
import os
import statistics

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.database.database import Database


class PeerReport:

    def __init__(self):

        self.db = Database()
        self.db.connect()

        self.output_folder = "output"

        os.makedirs(self.output_folder, exist_ok=True)

        self.workbook = Workbook()

        self.workbook.remove(self.workbook.active)

        # ========================================================
        # COLOURS
        # ========================================================

        self.green = PatternFill(fill_type="solid", start_color="90EE90")

        self.yellow = PatternFill(fill_type="solid", start_color="FFF59D")

        self.red = PatternFill(fill_type="solid", start_color="FFC7CE")

        self.gold = PatternFill(fill_type="solid", start_color="FFD966")

        self.header_fill = PatternFill(fill_type="solid", start_color="4472C4")

        self.header_font = Font(bold=True, color="FFFFFF")

        self.bold_font = Font(bold=True)

        # ========================================================
        # EXACT 20 METRICS FOR DAY 20
        # ========================================================

        self.metrics = [
            "Revenue Growth",
            "Profit Growth",
            "Net Profit Margin",
            "Operating Profit Margin",
            "ROE",
            "ROCE",
            "Debt To Equity",
            "Interest Coverage",
            "Asset Turnover",
            "Net Debt",
            "Free Cash Flow",
            "CFO Quality",
            "Capex Intensity",
            "Revenue CAGR 5Y",
            "Profit CAGR 5Y",
            "Quality Score",
            "Growth Score",
            "Value Score",
            "Risk Score",
            "Final Score",
        ]

    # ============================================================
    # LOAD PEER GROUPS
    # ============================================================

    def load_peer_groups(self):

        peer_groups = self.db.get_peer_groups()

        grouped = {}

        for group_name, company_id, benchmark in peer_groups:

            if group_name not in grouped:

                grouped[group_name] = []

            grouped[group_name].append(
                {"company_id": company_id, "benchmark": benchmark}
            )

        return grouped

    # ============================================================
    # COMPANY NAME
    # ============================================================

    def get_company_name(self, company_id):

        query = """
        SELECT company_name
        FROM companies
        WHERE id = ?
        """

        try:

            rows = self.db.execute_query(query, (company_id,))

            if rows and rows[0][0]:

                return rows[0][0]

        except Exception as e:

            print(f"Could not load company name for " f"{company_id}: {e}")

        return company_id

    # ============================================================
    # FINANCIAL RATIOS
    # ============================================================

    def get_financial_ratios(self, company_id):

        query = """
        SELECT
            revenue_growth,
            profit_growth,
            net_profit_margin,
            operating_profit_margin,
            roe,
            roce,
            debt_to_equity,
            interest_coverage,
            asset_turnover,
            net_debt,
            free_cash_flow,
            cfo_quality,
            capex_intensity
        FROM financial_ratios
        WHERE company_id = ?
        """

        try:

            rows = self.db.execute_query(query, (company_id,))

            if rows:

                row = rows[0]

                return {
                    "Revenue Growth": row[0],
                    "Profit Growth": row[1],
                    "Net Profit Margin": row[2],
                    "Operating Profit Margin": row[3],
                    "ROE": row[4],
                    "ROCE": row[5],
                    "Debt To Equity": row[6],
                    "Interest Coverage": row[7],
                    "Asset Turnover": row[8],
                    "Net Debt": row[9],
                    "Free Cash Flow": row[10],
                    "CFO Quality": row[11],
                    "Capex Intensity": row[12],
                }

        except Exception as e:

            print(f"Could not load financial ratios for " f"{company_id}: {e}")

        return {}

    # ============================================================
    # COMPANY RANKING SCORES
    # ============================================================

    def get_ranking_scores(self, company_id):

        query = """
        SELECT
            quality_score,
            growth_score,
            value_score,
            risk_score,
            final_score
        FROM company_rankings
        WHERE company_id = ?
        """

        try:

            rows = self.db.execute_query(query, (company_id,))

            if rows:

                row = rows[0]

                return {
                    "Quality Score": row[0],
                    "Growth Score": row[1],
                    "Value Score": row[2],
                    "Risk Score": row[3],
                    "Final Score": row[4],
                }

        except Exception as e:

            print(f"Could not load ranking scores for " f"{company_id}: {e}")

        return {}

    # ============================================================
    # CONVERT YEAR TO NUMBER
    # ============================================================

    def year_number(self, year):

        try:

            text = str(year).strip()

            digits = ""

            for character in text:

                if character.isdigit():

                    digits += character

            if len(digits) >= 4:

                return int(digits[:4])

        except Exception:

            pass

        return 0

    # ============================================================
    # CALCULATE 5-YEAR CAGR
    #
    # historical_financial_ratios contains annual
    # revenue_growth and profit_growth values.
    #
    # We compound the latest five annual growth rates.
    # ============================================================

    def calculate_cagr(self, company_id, column_name):

        allowed_columns = {"revenue_growth", "profit_growth"}

        if column_name not in allowed_columns:

            return None

        query = f"""
        SELECT
            year,
            {column_name}
        FROM historical_financial_ratios
        WHERE company_id = ?
        ORDER BY year
        """

        try:

            rows = self.db.execute_query(query, (company_id,))

            usable = []

            for row in rows:

                year = row[0]
                growth = row[1]

                if growth is None:

                    continue

                try:

                    growth = float(growth)

                except Exception:

                    continue

                usable.append((self.year_number(year), growth))

            # Sort by year

            usable.sort(key=lambda x: x[0])

            # Latest five years

            usable = usable[-5:]

            if len(usable) < 5:

                return None

            product = 1.0

            for _, growth in usable:

                factor = 1 + (growth / 100)

                # CAGR cannot be calculated if
                # an annual factor is zero/negative.

                if factor <= 0:

                    return None

                product *= factor

            years = len(usable)

            cagr = (product ** (1 / years) - 1) * 100

            return round(cagr, 2)

        except Exception as e:

            print(f"Could not calculate CAGR for " f"{company_id}: {e}")

            return None

    # ============================================================
    # GET ALL 20 METRICS
    # ============================================================

    def get_all_metrics(self, company_id):

        financial = self.get_financial_ratios(company_id)

        rankings = self.get_ranking_scores(company_id)

        revenue_cagr = self.calculate_cagr(company_id, "revenue_growth")

        profit_cagr = self.calculate_cagr(company_id, "profit_growth")

        data = {}

        # Financial ratios

        data.update(financial)

        # CAGR

        data["Revenue CAGR 5Y"] = revenue_cagr

        data["Profit CAGR 5Y"] = profit_cagr

        # Ranking scores

        data.update(rankings)

        return data

    # ============================================================
    # PERCENTILE
    # ============================================================

    def calculate_percentile(self, value, values):

        if value is None:

            return None

        clean_values = []

        for item in values:

            if item is None:

                continue

            try:

                number = float(item)

                if not math.isnan(number):

                    clean_values.append(number)

            except Exception:

                continue

        if not clean_values:

            return None

        try:

            value = float(value)

        except Exception:

            return None

        count = sum(1 for item in clean_values if item <= value)

        percentile = (count / len(clean_values)) * 100

        return round(percentile, 2)

    # ============================================================
    # APPLY PERCENTILE COLOUR
    # ============================================================

    def apply_percentile_colour(self, cell, percentile):

        if percentile is None:

            return

        if percentile >= 75:

            cell.fill = self.green

        elif percentile > 25:

            cell.fill = self.yellow

        else:

            cell.fill = self.red

    # ============================================================
    # BENCHMARK CHECK
    # ============================================================

    def is_benchmark(self, benchmark):

        return benchmark in (True, 1, "1", "true", "True", "TRUE", "yes", "Yes", "YES")

    # ============================================================
    # CREATE GROUP SHEET
    # ============================================================

    def create_group_sheet(self, group_name, companies):

        sheet = self.workbook.create_sheet(title=group_name[:31])

        # ========================================================
        # HEADERS
        # ========================================================

        headers = ["Company ID", "Company Name"]

        headers.extend(self.metrics)

        headers.extend([f"{metric} Percentile" for metric in self.metrics])

        # ========================================================
        # HEADER ROW
        # ========================================================

        for column, header in enumerate(headers, start=1):

            cell = sheet.cell(row=1, column=column)

            cell.value = header

            cell.fill = self.header_fill

            cell.font = self.header_font

            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

        # ========================================================
        # LOAD COMPANY DATA
        # ========================================================

        records = []

        for company in companies:

            company_id = company["company_id"]

            benchmark = company["benchmark"]

            metrics = self.get_all_metrics(company_id)

            records.append(
                {
                    "company_id": company_id,
                    "company_name": self.get_company_name(company_id),
                    "benchmark": benchmark,
                    "metrics": metrics,
                }
            )

        # ========================================================
        # CALCULATE PEER VALUES
        # ========================================================

        metric_values = {}

        for metric in self.metrics:

            metric_values[metric] = [
                record["metrics"].get(metric) for record in records
            ]

        # ========================================================
        # WRITE DATA
        # ========================================================

        current_row = 2

        percentile_start = 3 + len(self.metrics)

        for record in records:

            company_id = record["company_id"]

            company_name = record["company_name"]

            benchmark = record["benchmark"]

            metrics = record["metrics"]

            # ----------------------------------------------------
            # COMPANY ID
            # ----------------------------------------------------

            sheet.cell(row=current_row, column=1).value = company_id

            # ----------------------------------------------------
            # COMPANY NAME
            # ----------------------------------------------------

            sheet.cell(row=current_row, column=2).value = company_name

            # ----------------------------------------------------
            # METRIC VALUES
            # ----------------------------------------------------

            for index, metric in enumerate(self.metrics, start=3):

                value = metrics.get(metric)

                cell = sheet.cell(row=current_row, column=index)

                cell.value = value

                if isinstance(value, (int, float)):

                    cell.number_format = "0.00"

            # ----------------------------------------------------
            # PERCENTILE VALUES
            # ----------------------------------------------------

            for index, metric in enumerate(self.metrics, start=percentile_start):

                value = metrics.get(metric)

                percentile = self.calculate_percentile(value, metric_values[metric])

                cell = sheet.cell(row=current_row, column=index)

                cell.value = percentile

                if percentile is not None:

                    cell.number_format = "0.00"

                self.apply_percentile_colour(cell, percentile)

            # ----------------------------------------------------
            # BENCHMARK ROW
            #
            # Gold is applied to the company/name/metric cells.
            # Percentile cells keep their percentile colour.
            # ----------------------------------------------------

            if self.is_benchmark(benchmark):

                for column in range(1, percentile_start):

                    cell = sheet.cell(row=current_row, column=column)

                    cell.fill = self.gold

                    cell.font = self.bold_font

            current_row += 1

        # ========================================================
        # MEDIAN SUMMARY ROW
        # ========================================================

        median_row = current_row + 1

        sheet.cell(row=median_row, column=1).value = "PEER GROUP MEDIAN"

        sheet.cell(row=median_row, column=1).font = self.bold_font

        # --------------------------------------------------------
        # METRIC MEDIANS
        # --------------------------------------------------------

        for index, metric in enumerate(self.metrics, start=3):

            values = []

            for record in records:

                value = record["metrics"].get(metric)

                if value is not None:

                    try:

                        values.append(float(value))

                    except Exception:

                        pass

            if values:

                median_value = statistics.median(values)

                sheet.cell(row=median_row, column=index).value = round(median_value, 2)

        # --------------------------------------------------------
        # PERCENTILE MEDIANS
        # --------------------------------------------------------

        for index, metric in enumerate(self.metrics, start=percentile_start):

            percentile_values = []

            for record in records:

                value = record["metrics"].get(metric)

                percentile = self.calculate_percentile(value, metric_values[metric])

                if percentile is not None:

                    percentile_values.append(percentile)

            if percentile_values:

                median_percentile = statistics.median(percentile_values)

                cell = sheet.cell(row=median_row, column=index)

                cell.value = round(median_percentile, 2)

                self.apply_percentile_colour(cell, median_percentile)

        # ========================================================
        # SUMMARY ROW GOLD
        # ========================================================

        for column in range(1, percentile_start):

            cell = sheet.cell(row=median_row, column=column)

            cell.fill = self.gold
            cell.font = self.bold_font

        # ========================================================
        # FORMATTING
        # ========================================================

        sheet.freeze_panes = "C2"

        sheet.row_dimensions[1].height = 35

        last_column = len(headers)

        last_column_letter = get_column_letter(last_column)

        sheet.auto_filter.ref = f"A1:{last_column_letter}" f"{current_row - 1}"

        # Column widths

        sheet.column_dimensions["A"].width = 18

        sheet.column_dimensions["B"].width = 28

        for column in range(3, last_column + 1):

            letter = get_column_letter(column)

            sheet.column_dimensions[letter].width = 18

        print(f"Completed sheet: {group_name}")

    # ============================================================
    # CREATE ALL 11 SHEETS
    # ============================================================

    def create_sheets(self):

        groups = self.load_peer_groups()

        for group_name, companies in groups.items():

            self.create_group_sheet(group_name, companies)

        print(f"{len(groups)} Sheets Created")

    # ============================================================
    # SAVE
    # ============================================================

    def save(self):

        filename = os.path.join(self.output_folder, "peer_comparison.xlsx")

        self.workbook.save(filename)

        print("=" * 60)

        print("Excel Report Saved")

        print(filename)

        print("=" * 60)

    # ============================================================
    # CLOSE DATABASE
    # ============================================================

    def close(self):

        self.db.close()
