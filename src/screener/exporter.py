import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ScreenerExporter:

    OUTPUT_FILE = "output/screener_output.xlsx"

    PRESETS = {
        "quality_compounder": {
            "sheet": "Quality Compounder",
            "filters": {
                "roe": ("min", 15),
                "debt_to_equity": ("max", 1.0),
                "free_cash_flow": ("min", 0),
                "revenue_cagr_5yr": ("min", 10),
            },
        },
        "value_pick": {
            "sheet": "Value Pick",
            "filters": {
                "pe_ratio": ("max", 20),
                "pb_ratio": ("max", 3.0),
                "debt_to_equity": ("max", 2.0),
                "dividend_yield": ("min", 1),
            },
        },
        "growth_accelerator": {
            "sheet": "Growth Accelerator",
            "filters": {
                "pat_cagr_5yr": ("min", 20),
                "revenue_cagr_5yr": ("min", 15),
                "debt_to_equity": ("max", 2.0),
            },
        },
        "dividend_champion": {
            "sheet": "Dividend Champion",
            "filters": {
                "dividend_yield": ("min", 2),
                "dividend_payout": ("max", 80),
                "free_cash_flow": ("min", 0),
            },
        },
        "debt_free_bluechip": {
            "sheet": "Debt-Free Blue Chip",
            "filters": {
                "debt_to_equity": ("max", 0),
                "roe": ("min", 12),
                "sales": ("min", 5000),
            },
        },
        "turnaround_watch": {
            "sheet": "Turnaround Watch",
            "filters": {
                "revenue_cagr_3yr": ("min", 10),
                "free_cash_flow": ("min", 0),
                "de_declining": ("equals", True),
            },
        },
    }

    # Exactly 20 KPI columns required for Day 17.
    KPI_COLUMNS = [
        "roe",
        "roce",
        "net_profit_margin",
        "operating_profit_margin",
        "free_cash_flow",
        "fcf_cagr",
        "cfo_quality",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "eps_cagr_5yr",
        "debt_to_equity",
        "interest_coverage",
        "pe_ratio",
        "pb_ratio",
        "dividend_yield",
        "dividend_payout",
        "market_cap",
        "sales",
        "net_profit",
        "composite_quality_score",
    ]

    def __init__(self, engine):
        self.engine = engine

    # ==========================================================
    # APPLY PRESET THRESHOLD COLOUR
    # ==========================================================

    def _apply_threshold_colors(self, worksheet, df, preset_key):

        preset = self.PRESETS[preset_key]

        filters = preset["filters"]

        green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")

        red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

        green_font = Font(color="006100")

        red_font = Font(color="9C0006")

        # Map DataFrame column -> Excel column
        excel_columns = {}

        for column_number, column_name in enumerate(df.columns, start=1):
            excel_columns[column_name] = column_number

        for metric, rule in filters.items():

            if metric not in excel_columns:
                continue

            operation, threshold = rule

            excel_column = excel_columns[metric]

            for row_number in range(2, len(df) + 2):

                cell = worksheet.cell(row=row_number, column=excel_column)

                value = cell.value

                if value is None:
                    continue

                try:

                    if operation == "min":

                        passed = float(value) >= float(threshold)

                    elif operation == "max":

                        passed = float(value) <= float(threshold)

                    elif operation == "equals":

                        passed = bool(value) == bool(threshold)

                    else:

                        passed = False

                except (ValueError, TypeError):

                    passed = False

                if passed:

                    cell.fill = green_fill
                    cell.font = green_font

                else:

                    cell.fill = red_fill
                    cell.font = red_font

    # ==========================================================
    # FORMAT WORKSHEET
    # ==========================================================

    def _format_worksheet(self, worksheet):

        # Header
        for cell in worksheet[1]:

            cell.font = Font(bold=True)

            cell.alignment = Alignment(horizontal="center", vertical="center")

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = worksheet.dimensions

        # Column widths
        for column_cells in worksheet.columns:

            column_letter = get_column_letter(column_cells[0].column)

            max_length = 0

            for cell in column_cells:

                if cell.value is None:
                    continue

                value_length = len(str(cell.value))

                max_length = max(max_length, value_length)

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12), 28
            )

        # Number formatting
        for row in worksheet.iter_rows(min_row=2):

            for cell in row:

                if isinstance(cell.value, (int, float)):

                    cell.number_format = "0.00"

    # ==========================================================
    # PREPARE OUTPUT DATA
    # ==========================================================

    def _prepare_dataframe(self, preset_key):

        df = self.engine.filter_data(preset_key)

        if df is None:

            return pd.DataFrame()

        df = df.copy()

        # Ensure composite score exists.
        if "composite_quality_score" not in df.columns:

            df = self.engine.calculate_composite_score(df)

        # Sort by composite score.
        if "composite_quality_score" in df.columns:

            df = df.sort_values(by="composite_quality_score", ascending=False)

        # Company ID is needed for identification.
        output_columns = ["company_id", "broad_sector"]

        # Add the 20 KPI columns.
        for column in self.KPI_COLUMNS:

            if column in df.columns:

                output_columns.append(column)

            else:

                df[column] = None
                output_columns.append(column)

        # Remove duplicates while preserving order.
        output_columns = list(dict.fromkeys(output_columns))

        df = df[output_columns]

        return df.reset_index(drop=True)

    # ==========================================================
    # EXPORT ALL PRESETS
    # ==========================================================

    def export(self):

        output_directory = os.path.dirname(self.OUTPUT_FILE)

        os.makedirs(output_directory, exist_ok=True)

        # Remove old workbook if it exists.
        if os.path.exists(self.OUTPUT_FILE):

            os.remove(self.OUTPUT_FILE)

        exported = {}

        with pd.ExcelWriter(self.OUTPUT_FILE, engine="openpyxl") as writer:

            for preset_key, preset_info in self.PRESETS.items():

                print(f"Exporting: {preset_info['sheet']}")

                df = self._prepare_dataframe(preset_key)

                if df.empty:

                    print(f"WARNING: {preset_key} returned no data")

                    continue

                df.to_excel(writer, sheet_name=preset_info["sheet"], index=False)

                exported[preset_key] = len(df)

        # Re-open workbook for styling.
        workbook = load_workbook(self.OUTPUT_FILE)

        for preset_key, preset_info in self.PRESETS.items():

            sheet_name = preset_info["sheet"]

            if sheet_name not in workbook.sheetnames:
                continue

            worksheet = workbook[sheet_name]

            df = self._prepare_dataframe(preset_key)

            self._apply_threshold_colors(worksheet, df, preset_key)

            self._format_worksheet(worksheet)

        workbook.save(self.OUTPUT_FILE)

        print()
        print("=" * 70)
        print("DAY 17 EXCEL EXPORT COMPLETE")
        print("=" * 70)

        for preset_key, count in exported.items():

            print(f"{preset_key:<25} {count:>3} companies")

        print()
        print("OUTPUT:", os.path.abspath(self.OUTPUT_FILE))

        return self.OUTPUT_FILE
