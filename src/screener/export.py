import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ScreenerExporter:

    def __init__(self):

        self.output_folder = "output"

        os.makedirs(self.output_folder, exist_ok=True)

    def export(self, sheets):

        workbook = Workbook()

        workbook.remove(workbook.active)

        green = PatternFill(fill_type="solid", start_color="90EE90")

        red = PatternFill(fill_type="solid", start_color="FFC7CE")

        blue = PatternFill(fill_type="solid", start_color="4F81BD")

        header_font = Font(bold=True, color="FFFFFF")

        center = Alignment(horizontal="center", vertical="center")

        for sheet_name, df in sheets.items():

            ws = workbook.create_sheet(title=sheet_name)

            # Header
            ws.append(list(df.columns))

            for cell in ws[1]:

                cell.fill = blue
                cell.font = header_font
                cell.alignment = center

            # Data
            for row in df.itertuples(index=False):

                ws.append(row)

            # Colour numeric cells
            for row in ws.iter_rows(min_row=2):

                values = {}

                for i, cell in enumerate(row):

                    header = ws.cell(row=1, column=i + 1).value

                    values[header] = cell.value

                    cell.alignment = center

                for i, cell in enumerate(row):

                    header = ws.cell(row=1, column=i + 1).value

                    if not isinstance(cell.value, (int, float)):
                        continue

                    passed = True

                    if sheet_name == "QUALITY_COMPOUNDER":

                        if header == "roe":
                            passed = cell.value >= 15

                        elif header == "debt_to_equity":
                            passed = cell.value <= 1

                        elif header == "revenue_growth":
                            passed = cell.value >= 10

                        elif header == "free_cash_flow":
                            passed = cell.value > 0

                    elif sheet_name == "VALUE_PICK":

                        if header == "pe_ratio":
                            passed = cell.value <= 35

                        elif header == "pb_ratio":
                            passed = cell.value <= 6

                        elif header == "debt_to_equity":
                            passed = cell.value <= 3

                        elif header == "dividend_yield":
                            passed = cell.value >= 0

                    elif sheet_name == "GROWTH_ACCELERATOR":

                        if header == "revenue_growth" or header == "profit_growth":
                            passed = cell.value >= 10

                        elif header == "debt_to_equity":
                            passed = cell.value <= 2

                    elif sheet_name == "DIVIDEND_CHAMPION":

                        if header == "dividend_yield":
                            passed = cell.value >= 2

                        elif header == "free_cash_flow":
                            passed = cell.value > 0

                    elif sheet_name == "DEBT_FREE_BLUECHIP":

                        if header == "debt_to_equity":
                            passed = cell.value == 0

                        elif header == "roe":
                            passed = cell.value >= 12

                        elif header == "sales":
                            passed = cell.value >= 5000

                    elif sheet_name == "TURNAROUND_WATCH":

                        if header == "revenue_growth":
                            passed = cell.value >= 10

                        elif header == "free_cash_flow":
                            passed = cell.value > 0

                    cell.fill = green if passed else red

            # Auto column width
            for column in ws.columns:

                max_length = 0

                column_letter = get_column_letter(column[0].column)

                for cell in column:

                    try:
                        length = len(str(cell.value))

                        max_length = max(max_length, length)

                    except Exception:
                        pass

                ws.column_dimensions[column_letter].width = max_length + 2

            # Freeze header
            ws.freeze_panes = "A2"

            # Enable filter
            ws.auto_filter.ref = ws.dimensions

        workbook.save(os.path.join(self.output_folder, "screener_output.xlsx"))

        print("=" * 70)
        print("screener_output.xlsx Generated Successfully")
        print("=" * 70)
